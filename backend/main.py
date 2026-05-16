from fastapi import FastAPI, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse

import os
import uuid
import tempfile
import shutil

import fitz
from docx import Document

app = FastAPI()

BASE_URL = "https://pdf-tools-backend-rvzt.onrender.com"

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

OUTPUT_DIR = "outputs"

os.makedirs(OUTPUT_DIR, exist_ok=True)


@app.get("/")
def home():
    return {"status": "Backend running"}


# =====================================================
# SEARCH + HIGHLIGHT
# =====================================================
@app.post("/api/search-highlight")
async def search_highlight(
    file: UploadFile = File(...),
    search_text: str = Form(...),
    highlight_color: str = Form("yellow"),
):

    temp_dir = tempfile.mkdtemp()

    try:

        # -----------------------------
        # SAVE INPUT PDF
        # -----------------------------
        input_path = os.path.join(
            temp_dir,
            "input.pdf"
        )

        with open(input_path, "wb") as f:
            shutil.copyfileobj(file.file, f)

        # -----------------------------
        # OPEN PDF
        # -----------------------------
        pdf = fitz.open(input_path)

        # -----------------------------
        # COLORS
        # -----------------------------
        color_map = {
            "yellow": (1, 1, 0),
            "red": (1, 0, 0),
            "green": (0, 1, 0),
            "blue": (0, 0, 1),
            "pink": (1, 0.4, 0.7),
            "orange": (1, 0.5, 0),
        }

        selected_color = color_map.get(
            highlight_color.lower(),
            (1, 1, 0)
        )

        output_id = str(uuid.uuid4())

        matched_pages = []

        # =================================================
        # PROCESS PDF
        # =================================================
        for page_num in range(len(pdf)):

            page = pdf[page_num]

            matches = []

            # -----------------------------
            # GET PAGE TEXT
            # -----------------------------
            page_text = page.get_text("text")

            if not page_text:
                continue

            # -----------------------------
            # CASE INSENSITIVE CHECK
            # -----------------------------
            if search_text.lower() in page_text.lower():

                matched_pages.append(page_num)

                # -----------------------------
                # WORD SEARCH
                # -----------------------------
                words = page.get_text("words")

                for word in words:

                    x0, y0, x1, y1, text, *_ = word

                    if search_text.lower() in text.lower():

                        rect = fitz.Rect(x0, y0, x1, y1)

                        matches.append(rect)

            # -----------------------------
            # ADD HIGHLIGHTS
            # -----------------------------
            for rect in matches:

                annot = page.add_highlight_annot(rect)

                annot.set_colors(stroke=selected_color)

                annot.update()

        # =================================================
        # SAVE FULL PDF
        # =================================================
        full_pdf_path = os.path.join(
            OUTPUT_DIR,
            f"full-{output_id}.pdf"
        )

        pdf.save(full_pdf_path)

        # =================================================
        # CREATE MATCHED PDF
        # =================================================
        matched_pdf = fitz.open()

        if matched_pages:

            for page_num in matched_pages:

                matched_pdf.insert_pdf(
                    pdf,
                    from_page=page_num,
                    to_page=page_num
                )

        else:

            matched_pdf.new_page()

        matched_pdf_path = os.path.join(
            OUTPUT_DIR,
            f"matched-{output_id}.pdf"
        )

        matched_pdf.save(matched_pdf_path)

        matched_pdf.close()

        pdf.close()

        # =================================================
        # RESPONSE
        # =================================================
        return {
            "success": True,
            "full_pdf_url": f"{BASE_URL}/download-full-pdf/{output_id}",
            "matched_pdf_url": f"{BASE_URL}/download-matched-pdf/{output_id}"
        }

    except Exception as e:

        return {
            "success": False,
            "error": str(e)
        }

    finally:

        shutil.rmtree(temp_dir, ignore_errors=True)


# =====================================================
# PDF TO WORD
# =====================================================
@app.post("/api/pdf-to-word")
async def pdf_to_word(
    file: UploadFile = File(...)
):

    temp_dir = tempfile.mkdtemp()

    try:

        input_path = os.path.join(
            temp_dir,
            "input.pdf"
        )

        with open(input_path, "wb") as f:
            shutil.copyfileobj(file.file, f)

        pdf = fitz.open(input_path)

        text = ""

        for page in pdf:

            text += page.get_text("text") + "\n"

        pdf.close()

        output_id = str(uuid.uuid4())

        output_path = os.path.join(
            OUTPUT_DIR,
            f"converted-{output_id}.docx"
        )

        word = Document()

        word.add_heading(
            "Converted PDF",
            level=1
        )

        if text.strip():

            for line in text.split("\n"):

                line = line.strip()

                if line:
                    word.add_paragraph(line)

        else:

            word.add_paragraph(
                "No extractable text found"
            )

        word.save(output_path)

        return {
            "success": True,
            "download_url": f"{BASE_URL}/download-docx/{output_id}"
        }

    except Exception as e:

        return {
            "success": False,
            "error": str(e)
        }

    finally:

        shutil.rmtree(temp_dir, ignore_errors=True)


# =====================================================
# DOWNLOAD FULL PDF
# =====================================================
@app.get("/download-full-pdf/{file_id}")
def download_full_pdf(file_id: str):

    path = os.path.join(
        OUTPUT_DIR,
        f"full-{file_id}.pdf"
    )

    if not os.path.exists(path):

        return {
            "error": "File not found"
        }

    return FileResponse(
        path,
        media_type="application/pdf",
        filename="highlighted-full.pdf"
    )


# =====================================================
# DOWNLOAD MATCHED PAGES PDF
# =====================================================
@app.get("/download-matched-pdf/{file_id}")
def download_matched_pdf(file_id: str):

    path = os.path.join(
        OUTPUT_DIR,
        f"matched-{file_id}.pdf"
    )

    if not os.path.exists(path):

        return {
            "error": "File not found"
        }

    return FileResponse(
        path,
        media_type="application/pdf",
        filename="highlighted-pages.pdf"
    )


# =====================================================
# DOWNLOAD DOCX
# =====================================================
@app.get("/download-docx/{file_id}")
def download_docx(file_id: str):

    path = os.path.join(
        OUTPUT_DIR,
        f"converted-{file_id}.docx"
    )

    if not os.path.exists(path):

        return {
            "error": "File not found"
        }

    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="converted.docx"
    )